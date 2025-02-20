import streamlit as st
import google.generativeai as genai
from PyPDF2 import PdfReader
import docx
import os
from dotenv import load_dotenv
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from datetime import datetime
import json
from translate import Translator
from fpdf import FPDF
import io
import base64

# Load environment variables
load_dotenv()

# Configure Gemini API
genai.configure(api_key=os.getenv('GOOGLE_API_KEY'))
model = genai.GenerativeModel('gemini-pro')

# Supported languages
LANGUAGES = {
    'English': 'en',
    'Spanish': 'es',
    'French': 'fr',
    'German': 'de',
    'Chinese': 'zh',
    'Japanese': 'ja'
}

def translate_text(text, target_lang):
    """Translate text to target language"""
    try:
        translator = Translator(to_lang=target_lang)
        chunks = [text[i:i+500] for i in range(0, len(text), 500)]
        translated_text = ""
        for chunk in chunks:
            translated_text += translator.translate(chunk) + " "
        return translated_text
    except Exception as e:
        st.error(f"Translation error: {str(e)}")
        return text

def extract_text_from_pdf(pdf_file):
    """Extract text from PDF file"""
    try:
        pdf_reader = PdfReader(pdf_file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        st.error(f"PDF extraction error: {str(e)}")
        return ""

def extract_text_from_docx(docx_file):
    """Extract text from DOCX file"""
    try:
        doc = docx.Document(docx_file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        st.error(f"DOCX extraction error: {str(e)}")
        return ""

def analyze_contract_score(text):
    """Calculate overall contract score and return structured data"""
    prompt = f"""
    Analyze this contract and provide a comprehensive score in the following EXACT JSON format:
    {{
        "overall_score": 85,
        "score_breakdown": {{
            "clarity": 90,
            "completeness": 85,
            "fairness": 80,
            "risk_management": 85,
            "compliance": 85
        }},
        "summary": "Brief summary of the score justification"
    }}
    Contract text: {text}
    Only return the JSON object, no other text.
    """
    
    try:
        response = model.generate_content(prompt)
        cleaned_text = response.text.strip()
        if cleaned_text.startswith("```json"):
            cleaned_text = cleaned_text[7:-3]
        return json.loads(cleaned_text)
    except Exception as e:
        return {
            "overall_score": 0,
            "score_breakdown": {
                "clarity": 0,
                "completeness": 0,
                "fairness": 0,
                "risk_management": 0,
                "compliance": 0
            },
            "summary": "Score calculation failed"
        }

def analyze_risks_and_opportunities(text):
    """Analyze contract risks and opportunities"""
    prompt = f"""
    Analyze this contract and provide risks and opportunities in the following EXACT JSON format:
    {{
        "risks": {{
            "legal_risks": {{
                "level": "High/Medium/Low",
                "description": "description",
                "potential_impact": "impact",
                "mitigation_suggestions": "suggestions"
            }},
            "financial_risks": {{
                "level": "High/Medium/Low",
                "description": "description",
                "potential_impact": "impact",
                "mitigation_suggestions": "suggestions"
            }},
            "operational_risks": {{
                "level": "High/Medium/Low",
                "description": "description",
                "potential_impact": "impact",
                "mitigation_suggestions": "suggestions"
            }}
        }},
        "opportunities": {{
            "business_opportunities": {{
                "level": "High/Medium/Low",
                "description": "description",
                "potential_value": "value",
                "action_items": "actions"
            }},
            "strategic_opportunities": {{
                "level": "High/Medium/Low",
                "description": "description",
                "potential_value": "value",
                "action_items": "actions"
            }},
            "operational_opportunities": {{
                "level": "High/Medium/Low",
                "description": "description",
                "potential_value": "value",
                "action_items": "actions"
            }}
        }}
    }}
    Contract text: {text}
    """
    
    try:
        response = model.generate_content(prompt)
        cleaned_text = response.text.strip()
        if cleaned_text.startswith("```json"):
            cleaned_text = cleaned_text[7:-3]
        return json.loads(cleaned_text)
    except Exception as e:
        return {"risks": {}, "opportunities": {}}

def analyze_key_clauses(text):
    """Analyze key clauses and provide recommendations"""
    prompt = f"""
    Analyze this contract and provide key clauses and recommendations in the following EXACT JSON format:
    {{
        "key_clauses": {{
            "critical_clauses": [
                {{
                    "title": "clause title",
                    "summary": "brief summary",
                    "importance": "High/Medium/Low"
                }}
            ],
            "notable_provisions": [
                {{
                    "title": "provision title",
                    "summary": "brief summary",
                    "impact": "description"
                }}
            ]
        }},
        "recommendations": [
            {{
                "category": "category name",
                "suggestion": "detailed suggestion",
                "priority": "High/Medium/Low"
            }}
        ]
    }}
    Contract text: {text}
    """
    
    try:
        response = model.generate_content(prompt)
        cleaned_text = response.text.strip()
        if cleaned_text.startswith("```json"):
            cleaned_text = cleaned_text[7:-3]
        return json.loads(cleaned_text)
    except Exception as e:
        return {"key_clauses": {"critical_clauses": [], "notable_provisions": []}, "recommendations": []}
    
def create_risk_opportunity_charts(analysis_data):
    """Create visualizations for risks and opportunities"""
    # Process risks and opportunities
    risk_levels = {"High": 3, "Medium": 2, "Low": 1}
    
    # Calculate percentages
    total_items = len(analysis_data["risks"]) + len(analysis_data["opportunities"])
    risk_percentage = (len(analysis_data["risks"]) / total_items) * 100
    opp_percentage = (len(analysis_data["opportunities"]) / total_items) * 100
    
    # Create pie chart
    pie_data = pd.DataFrame({
        'Category': ['Risks', 'Opportunities'],
        'Percentage': [risk_percentage, opp_percentage]
    })
    
    pie_fig = px.pie(
        pie_data,
        values='Percentage',
        names='Category',
        title='Distribution of Risks vs Opportunities',
        color_discrete_sequence=['#FF6B6B', '#4ECDC4']
    )
    
    # Prepare radar chart data
    risk_data = [
        {
            'category': k.replace('_', ' ').title(),
            'value': risk_levels[v['level']],
            'type': 'Risk'
        }
        for k, v in analysis_data['risks'].items()
    ]
    
    opp_data = [
        {
            'category': k.replace('_', ' ').title(),
            'value': risk_levels[v['level']],
            'type': 'Opportunity'
        }
        for k, v in analysis_data['opportunities'].items()
    ]
    
    # Create radar chart
    radar_fig = go.Figure()
    
    radar_fig.add_trace(go.Scatterpolar(
        r=[d['value'] for d in risk_data],
        theta=[d['category'] for d in risk_data],
        name='Risks',
        fill='toself',
        fillcolor='rgba(255, 107, 107, 0.5)'
    ))
    
    radar_fig.add_trace(go.Scatterpolar(
        r=[d['value'] for d in opp_data],
        theta=[d['category'] for d in opp_data],
        name='Opportunities',
        fill='toself',
        fillcolor='rgba(78, 205, 196, 0.5)'
    ))
    
    radar_fig.update_layout(
        polar=dict(radialaxis=dict(visible=True, range=[0, 3])),
        showlegend=True,
        title='Risks and Opportunities Assessment'
    )
    
    return pie_fig, radar_fig

def generate_pdf_report(contract_text, analysis_results, score_data, risk_opp_data, clause_data, figures):
    """Generate PDF report with analysis results and visualizations"""
    class PDF(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 15)
            self.cell(0, 10, 'Contract Analysis Report', 0, 1, 'C')
            self.ln(10)
            
        def footer(self):
            self.set_y(-15)
            self.set_font('Arial', 'I', 8)
            self.cell(0, 10, f'Generated on {datetime.now().strftime("%Y-%m-%d %H:%M:%S")} | Page {self.page_no()}', 0, 0, 'C')
    
    pdf = PDF()
    
    # Executive Summary
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Executive Summary', 0, 1)
    pdf.set_font('Arial', '', 11)
    pdf.multi_cell(0, 10, score_data["summary"])
    
    # Contract Score
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Contract Score Analysis', 0, 1)
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, f'Overall Score: {score_data["overall_score"]}/100', 0, 1)
    
    # Score Breakdown
    pdf.set_font('Arial', '', 11)
    for metric, score in score_data["score_breakdown"].items():
        pdf.cell(0, 10, f'{metric.title()}: {score}/100', 0, 1)
    
    # Risks and Opportunities
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Risks Analysis', 0, 1)
    
    for risk_type, details in risk_opp_data["risks"].items():
        pdf.set_font('Arial', 'B', 12)
        pdf.cell(0, 10, risk_type.replace("_", " ").title(), 0, 1)
        pdf.set_font('Arial', '', 11)
        pdf.multi_cell(0, 10, f'''
Level: {details["level"]}
Description: {details["description"]}
Impact: {details["potential_impact"]}
Mitigation: {details["mitigation_suggestions"]}
        ''')
    
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Opportunities Analysis', 0, 1)
    
    for opp_type, details in risk_opp_data["opportunities"].items():
        pdf.set_font('Arial', 'B', 12)
        pdf.cell(0, 10, opp_type.replace("_", " ").title(), 0, 1)
        pdf.set_font('Arial', '', 11)
        pdf.multi_cell(0, 10, f'''
Level: {details["level"]}
Description: {details["description"]}
Potential Value: {details["potential_value"]}
Action Items: {details["action_items"]}
        ''')
    
    # Key Clauses
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Key Clauses Analysis', 0, 1)
    
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, 'Critical Clauses', 0, 1)
    for clause in clause_data["key_clauses"]["critical_clauses"]:
        pdf.set_font('Arial', 'B', 11)
        pdf.cell(0, 10, clause["title"], 0, 1)
        pdf.set_font('Arial', '', 11)
        pdf.multi_cell(0, 10, f'''
Summary: {clause["summary"]}
Importance: {clause["importance"]}
        ''')
    
    # Recommendations
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Recommendations', 0, 1)
    
    for rec in clause_data["recommendations"]:
        pdf.set_font('Arial', 'B', 11)
        pdf.cell(0, 10, f'{rec["category"]} (Priority: {rec["priority"]})', 0, 1)
        pdf.set_font('Arial', '', 11)
        pdf.multi_cell(0, 10, rec["suggestion"])
    
    # Add visualizations
    pdf.add_page()
    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, 'Visual Analysis', 0, 1)
    
    for fig in figures:
        img_bytes = fig.to_image(format="png")
        pdf.image(io.BytesIO(img_bytes), x=10, w=190)
    
    return pdf.output(dest='S').encode('latin1')

def ask_question(text, question, context=None):
    """Enhanced Q&A function with context tracking"""
    try:
        prompt = f"""
        Based on this contract: {text}
        
        Previous context: {context if context else 'None'}
        
        Please answer this question: {question}
        
        Provide a clear, direct answer with specific references to relevant sections of the contract when applicable.
        """
        
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        st.error(f"Q&A error: {str(e)}")
        return "Failed to get answer. Please try again."

def main():
    st.set_page_config(page_title="Contract Analysis System", layout="wide")
    
    st.title("Advanced Contract Analysis System")
    st.write("Upload your contracts for AI-powered analysis")
    
    # Language selection
    selected_language = st.selectbox("Select Language", list(LANGUAGES.keys()))
    
    # Session state initialization
    if 'qa_history' not in st.session_state:
        st.session_state.qa_history = []
    
    # File upload
    uploaded_files = st.file_uploader(
        "Choose contract file(s)", 
        type=['pdf', 'docx'], 
        accept_multiple_files=True,
        help="Upload PDF or DOCX files for analysis"
    )
    
    if uploaded_files:
        contract_texts = []
        for file in uploaded_files:
            if file.type == "application/pdf":
                text = extract_text_from_pdf(file)
            else:
                text = extract_text_from_docx(file)
            contract_texts.append(text)
        
        # Create tabs for different analyses
        tabs = st.tabs(["Overview", "Risks & Opportunities", "Key Clauses", "Comparison", "Q&A", "Report"])
        
        with tabs[0]:
            st.header("Contract Overview")
            if st.button("Analyze Overview"):
                for i, text in enumerate(contract_texts):
                    st.subheader(f"Contract {i+1}")
                    with st.spinner("Analyzing..."):
                        # Get contract score
                        score_data = analyze_contract_score(text)
                        
                        # Create score visualization
                        col1, col2 = st.columns(2)
                        with col1:
                            st.metric("Overall Score", f"{score_data['overall_score']}/100")
                            
                            # Score breakdown
                            st.subheader("Score Breakdown")
                            for metric, score in score_data["score_breakdown"].items():
                                st.metric(metric.title(), f"{score}/100")
                        
                        with col2:
                            # Gauge chart for overall score
                            gauge_fig = go.Figure(go.Indicator(
                                mode="gauge+number",
                                value=score_data['overall_score'],
                                title={'text': "Contract Score"},
                                gauge={'axis': {'range': [0, 100]},
                                       'bar': {'color': "#4ECDC4"},
                                       'steps': [
                                           {'range': [0, 50], 'color': "#FF6B6B"},
                                           {'range': [50, 75], 'color': "#FFD93D"},
                                           {'range': [75, 100], 'color': "#4ECDC4"}
                                       ]}
                            ))
                            st.plotly_chart(gauge_fig)
                        
                        st.markdown("### Analysis Summary")
                        st.write(score_data["summary"])
        
        with tabs[1]:
            st.header("Risks & Opportunities Assessment")
            if st.button("Analyze Risks & Opportunities"):
                for i, text in enumerate(contract_texts):
                    st.subheader(f"Contract {i+1}")
                    with st.spinner("Analyzing risks and opportunities..."):
                        analysis_data = analyze_risks_and_opportunities(text)
                        
                        # Create visualizations
                        pie_fig, radar_fig = create_risk_opportunity_charts(analysis_data)
                        
                        col1, col2 = st.columns(2)
                        with col1:
                            st.plotly_chart(pie_fig, use_container_width=True)
                        with col2:
                            st.plotly_chart(radar_fig, use_container_width=True)
                        
                        # Display detailed analysis
                        col1, col2 = st.columns(2)
                        
                        with col1:
                            st.subheader("Risks")
                            for risk_type, details in analysis_data["risks"].items():
                                with st.expander(f"{risk_type.replace('_', ' ').title()} ({details['level']})"):
                                    st.markdown(f"**Description:** {details['description']}")
                                    st.markdown(f"**Impact:** {details['potential_impact']}")
                                    st.markdown(f"**Mitigation:** {details['mitigation_suggestions']}")
                        
                        with col2:
                            st.subheader("Opportunities")
                            for opp_type, details in analysis_data["opportunities"].items():
                                with st.expander(f"{opp_type.replace('_', ' ').title()} ({details['level']})"):
                                    st.markdown(f"**Description:** {details['description']}")
                                    st.markdown(f"**Value:** {details['potential_value']}")
                                    st.markdown(f"**Actions:** {details['action_items']}")
        
        with tabs[2]:
            st.header("Key Clauses Analysis")
            if st.button("Analyze Key Clauses"):
                for i, text in enumerate(contract_texts):
                    st.subheader(f"Contract {i+1}")
                    with st.spinner("Analyzing key clauses..."):
                        clause_data = analyze_key_clauses(text)
                        
                        # Display critical clauses
                        st.subheader("Critical Clauses")
                        for clause in clause_data["key_clauses"]["critical_clauses"]:
                            with st.expander(f"{clause['title']} (Importance: {clause['importance']})"):
                                st.write(clause["summary"])
                        
                        # Display recommendations
                        st.subheader("Recommendations")
                        for rec in clause_data["recommendations"]:
                            with st.expander(f"{rec['category']} (Priority: {rec['priority']})"):
                                st.write(rec["suggestion"])
        
        with tabs[3]:
            st.header("Contract Comparison")
            if len(contract_texts) >= 2:
                if st.button("Compare Contracts"):
                    with st.spinner("Comparing contracts..."):
                        # Compare risk profiles
                        scores = []
                        analyses = []
                        for text in contract_texts[:2]:  # Compare first two contracts
                            scores.append(analyze_contract_score(text))
                            analyses.append(analyze_risks_and_opportunities(text))
                        
                        col1, col2 = st.columns(2)
                        with col1:
                            st.subheader("Score Comparison")
                            score_comparison = pd.DataFrame({
                                'Metric': ['Overall'] + list(scores[0]['score_breakdown'].keys()),
                                'Contract 1': [scores[0]['overall_score']] + list(scores[0]['score_breakdown'].values()),
                                'Contract 2': [scores[1]['overall_score']] + list(scores[1]['score_breakdown'].values())
                            })
                            fig = px.bar(score_comparison, x='Metric', y=['Contract 1', 'Contract 2'],
                                       title="Score Comparison", barmode='group')
                            st.plotly_chart(fig)
                        
                        with col2:
                            st.subheader("Risk Profile Comparison")
                            # Create comparison visualization
                            risk_comparison = go.Figure()
                            for i, analysis in enumerate(analyses, 1):
                                risk_levels = [{'High': 3, 'Medium': 2, 'Low': 1}[risk['level']] 
                                             for risk in analysis['risks'].values()]
                                risk_comparison.add_trace(go.Scatterpolar(
                                    r=risk_levels,
                                    theta=list(analysis['risks'].keys()),
                                    name=f'Contract {i}'
                                ))
                            risk_comparison.update_layout(title="Risk Profile Comparison")
                            st.plotly_chart(risk_comparison)
            else:
                st.warning("Please upload at least two contracts for comparison")
        
        with tabs[4]:
            st.header("Question & Answer")
            selected_contract = st.selectbox(
                "Select Contract for Q&A",
                range(len(contract_texts)),
                format_func=lambda x: f"Contract {x+1}"
            )
            
            # Question templates
            question_templates = [
                "What are the key deadlines?",
                "What are the main financial obligations?",
                "What are the termination conditions?",
                "What are the liability clauses?",
                "Custom question"
            ]
            
            template_choice = st.selectbox("Choose question type:", question_templates)
            
            if template_choice == "Custom question":
                question = st.text_input("Enter your question:")
            else:
                question = template_choice
            
            if question and st.button("Ask Question"):
                with st.spinner("Finding answer..."):
                    context = "; ".join([f"Q: {q} A: {a}" for q, a in st.session_state.qa_history[-3:]])
                    answer = ask_question(contract_texts[selected_contract], question, context)
                    
                    if selected_language != 'English':
                        answer = translate_text(answer, LANGUAGES[selected_language])
                    
                    st.session_state.qa_history.append((question, answer))
                    st.markdown("### Answer:")
                    st.markdown(answer)
            
            if st.session_state.qa_history:
                with st.expander("Question & Answer History"):
                    for q, a in st.session_state.qa_history:
                        st.markdown(f"**Q:** {q}")
                        st.markdown(f"**A:** {a}")
                        st.markdown("---")
        
        with tabs[5]:
            st.header("Generate Report")
            if st.button("Generate Complete Report"):
                for i, text in enumerate(contract_texts):
                    with st.spinner(f"Generating report for Contract {i+1}..."):
                        # Gather all analysis data
                        score_data = analyze_contract_score(text)
                        risk_opp_data = analyze_risks_and_opportunities(text)
                        clause_data = analyze_key_clauses(text)
                        
                        # Create visualizations
                        pie_fig, radar_fig = create_risk_opportunity_charts(risk_opp_data)
                        
                        # Generate PDF report
                        pdf_bytes = generate_pdf_report(
                            text,
                            score_data,
                            risk_opp_data,
                            clause_data,
                            [pie_fig, radar_fig]
                        )
                        
                        # Create download button
                        st.download_button(
                            label=f"Download Report (Contract {i+1})",
                            data=pdf_bytes,
                            file_name=f"contract_analysis_report_{i+1}.pdf",
                            mime="application/pdf"
                        )
                        
                        # Display report preview
                        st.markdown("### Report Preview")
                        st.markdown(generate_pdf_report(
                            score_data,
                            risk_opp_data,
                            clause_data
                        ))

if __name__ == "__main__":
    main()       